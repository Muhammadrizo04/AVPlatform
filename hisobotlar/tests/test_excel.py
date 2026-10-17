import io

from openpyxl import load_workbook

from hisobotlar.excel import jadval_xlsx


def test_formula_inyeksiyasi_matn_sifatida_saqlanadi():
    baytlar = jadval_xlsx(
        ["=SARLAVHA()", "Nomi"],
        [["=1+1", "@SUM(A1)"], ["+5", "-3"], ["\tx", "oddiy"], [7, "=cmd|' /C calc'!A0"]],
    )
    varaq = load_workbook(io.BytesIO(baytlar)).active
    for qator in varaq.iter_rows():
        for yacheyka in qator:
            assert yacheyka.data_type != "f", yacheyka.coordinate
    assert varaq["A1"].value == "=SARLAVHA()"
    assert varaq["A2"].value == "=1+1"
    assert varaq["B2"].value == "@SUM(A1)"
    assert varaq["A5"].value == 7
    assert varaq["B5"].value == "=cmd|' /C calc'!A0"
