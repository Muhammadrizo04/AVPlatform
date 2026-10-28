import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from hisobotlar import servis
from onec.models import AVTuri
from vositalar.amortizatsiya import dona_narxi, oy_qosh
from vositalar.models import AVKartochka, Bolim
from vositalar.tests.test_bayroqlar import BUGUN, av_yarat, qator, xodim


@pytest.fixture
def filial_a(filial_yarat):
    return filial_yarat("A", "Filial A")


@pytest.fixture
def filial_b(filial_yarat):
    return filial_yarat("B", "Filial B")


@pytest.fixture
def hisobchi(foydalanuvchi):
    return foydalanuvchi("hisobchi")


def kartochka(av, ishga_tushgan, muddat):
    return AVKartochka.objects.create(av=av, ishga_tushgan_sana=ishga_tushgan, xizmat_muddati_oy=muddat)


def tur(nomi):
    return AVTuri.objects.create(ref_1c=uuid.uuid4(), nomi=nomi)


def xlsx_qatorlar(javob):
    varaq = load_workbook(io.BytesIO(javob.content)).active
    return [[c.value for c in qator] for qator in varaq.iter_rows()]


def test_kesim_filial_jamilari(hisobchi, filial_a, filial_b):
    av_yarat(filial_a, soni="2")  # summa 100, kartochkasiz -> qoldiq = summa
    eskirgan = av_yarat(filial_a, soni="3")  # 12/24 oy -> 50%
    kartochka(eskirgan, oy_qosh(BUGUN, -12), 24)
    av_yarat(filial_a, soni="0")  # arxiv — hisobga olinmaydi
    av_yarat(filial_b, soni="1")

    qatorlar = {q["nomi"]: q for q in servis.kesim(hisobchi, "filial", BUGUN)}

    assert qatorlar["Filial A"] == {
        "nomi": "Filial A", "soni": Decimal("5"), "boshlangich": Decimal("200"), "qoldiq": Decimal("150.00"),
    }
    assert qatorlar["Filial B"]["boshlangich"] == Decimal("100")


def test_kesim_bolim_guruh_ikki_qatorga_bolinadi(hisobchi, filial_a):
    guruh = av_yarat(filial_a, soni="20")  # summa 100
    x = Bolim.objects.create(filial=filial_a, nomi="Kassa")
    y = Bolim.objects.create(filial=filial_a, nomi="Ombor", ota=x)
    qator(guruh, filial_a, "12", bolim=x)
    qator(guruh, filial_a, "5", bolim=y)
    qator(guruh, filial_a, "3")  # taqsimlanmagan

    qatorlar = {q["nomi"]: q for q in servis.kesim(hisobchi, "bolim", BUGUN)}

    assert set(qatorlar) == {"Filial A / Kassa", "Filial A / Kassa / Ombor", "Bo'limsiz"}
    assert qatorlar["Filial A / Kassa"]["soni"] == Decimal("12")
    assert qatorlar["Filial A / Kassa"]["boshlangich"] == Decimal("60")
    assert qatorlar["Filial A / Kassa / Ombor"]["boshlangich"] == Decimal("25")
    assert qatorlar["Bo'limsiz"]["boshlangich"] == Decimal("15")
    assert sum(q["boshlangich"] for q in qatorlar.values()) == Decimal("100")  # 1C summasi


def test_kesim_bolim_eskirish(hisobchi, filial_a):
    av = av_yarat(filial_a, soni="4")
    kartochka(av, oy_qosh(BUGUN, -12), 24)
    qator(av, filial_a, "4")
    (q,) = servis.kesim(hisobchi, "bolim", BUGUN)
    assert q["nomi"] == "Bo'limsiz"
    assert (q["boshlangich"], q["qoldiq"]) == (Decimal("100"), Decimal("50.00"))


def test_kesim_tur(hisobchi, filial_a):
    av_yarat(filial_a, turi=tur("Kompyuter"))
    av_yarat(filial_a)
    qatorlar = {q["nomi"]: q for q in servis.kesim(hisobchi, "tur", BUGUN)}
    assert set(qatorlar) == {"Kompyuter", "Turi yo'q"}


def test_kesim_notogri_guruh(hisobchi):
    with pytest.raises(ValueError):
        servis.kesim(hisobchi, "xato", BUGUN)


def test_mol_qaydnoma_qiymat(hisobchi, filial_a):
    av = av_yarat(filial_a, soni="4")  # summa 100 -> dona 25
    ali = xodim()
    bolim = Bolim.objects.create(filial=filial_a, nomi="Kassa")
    qator(av, filial_a, "3", bolim=bolim, mol=ali)
    qator(av, filial_a, "1")  # MOL'siz — kirmaydi

    (q,) = servis.mol_qaydnoma(hisobchi)

    assert q["qiymat"] == Decimal("3") * dona_narxi(Decimal("100"), Decimal("4")) == Decimal("75")
    assert (q["xodim"], q["av"], q["soni"], q["bolim"]) == (ali, av, Decimal("3"), "Kassa")
    assert servis.mol_qaydnoma(hisobchi, ali) == [q]
    assert servis.mol_qaydnoma(hisobchi, xodim()) == []


def test_muddat_chegaralari(hisobchi, filial_a):
    tugagan = av_yarat(filial_a)
    kartochka(tugagan, oy_qosh(BUGUN, -12), 12)  # tugash = bugun
    olti = av_yarat(filial_a)
    kartochka(olti, BUGUN, 6)  # tugash = bugun + 6 oy
    tashqarida = av_yarat(filial_a)
    kartochka(tashqarida, date(2026, 10, 18), 6)  # bugun + 6 oy + 1 kun

    t, o = servis.muddat(hisobchi, BUGUN)

    assert [av.pk for av in t] == [tugagan.pk]
    assert [av.pk for av in o] == [olti.pk]


def test_tozalash_faqat_1c_bayroqlari(client_rol, filial_a):
    bosh = av_yarat(filial_a, inv="")
    rekvizitsiz = av_yarat(None, qoldiq_filial=filial_a)
    javob = client_rol("hisobchi").get(reverse("hisobotlar:tozalash"))
    assert javob.status_code == 200
    bolimlar = {b["kalit"]: b for b in javob.context["bolimlar"]}
    assert list(bolimlar) == ["inv_bosh", "inv_takror", "filial_rekvizit_bosh", "mos_emas"]
    assert bosh.pk in {av.pk for av in bolimlar["inv_bosh"]["avlar"]}
    assert rekvizitsiz.pk in {av.pk for av in bolimlar["filial_rekvizit_bosh"]["avlar"]}


def test_huquqlar(client_rol):
    hr = client_rol("hr")
    assert hr.get(reverse("hisobotlar:tozalash")).status_code == 403
    assert hr.get(reverse("hisobotlar:tozalash"), {"format": "xlsx"}).status_code == 403
    for nom in ("kesim", "mol_qaydnoma", "muddat"):
        assert hr.get(reverse(f"hisobotlar:{nom}")).status_code == 200
    assert Client().get(reverse("hisobotlar:kesim")).status_code == 302


def test_notogri_parametr_500_bermaydi(client_rol):
    c = client_rol("hisobchi")
    assert c.get(reverse("hisobotlar:kesim"), {"guruh": "xato"}).status_code == 404
    for qiymat in ("abc", "²", "9" * 40):
        assert c.get(reverse("hisobotlar:mol_qaydnoma"), {"xodim": qiymat}).status_code == 404
    assert c.get(reverse("hisobotlar:mol_qaydnoma"), {"xodim": "999999"}).status_code == 404


def test_hisobot_excel_filial_cheklovi(foydalanuvchi, filial_a, filial_b):
    av_a = av_yarat(filial_a, inv="")
    av_b = av_yarat(filial_b, inv="")
    ali = xodim()
    bolim_b = Bolim.objects.create(filial=filial_b, nomi="Bolim B")
    qator(av_a, filial_a, "1", mol=ali)
    qator(av_b, filial_b, "1", mol=ali, bolim=bolim_b)
    kartochka(av_a, oy_qosh(BUGUN, -24), 12)
    kartochka(av_b, oy_qosh(BUGUN, -24), 12)
    client = Client()
    client.force_login(foydalanuvchi("hisobchi", [filial_a]))
    sahifalar = [
        ("kesim", {"guruh": "filial"}),
        ("kesim", {"guruh": "bolim"}),
        ("kesim", {"guruh": "tur"}),
        ("tozalash", {}),
        ("mol_qaydnoma", {}),
        ("muddat", {}),
    ]
    for nom, params in sahifalar:
        javob = client.get(reverse(f"hisobotlar:{nom}"), {**params, "format": "xlsx"})
        assert javob.status_code == 200, (nom, params)
        assert "spreadsheetml" in javob["Content-Type"]
        matn = str(xlsx_qatorlar(javob))
        assert av_b.nomi not in matn and "Filial B" not in matn and "Bolim B" not in matn, (nom, params)
        if nom in ("tozalash", "mol_qaydnoma", "muddat"):
            assert av_a.nomi in matn, nom


def test_hisobotlar_sorovlar_soni_qatorlarga_bogliq_emas(hisobchi, filial_a, django_assert_max_num_queries):
    ali = xodim()
    ota = Bolim.objects.create(filial=filial_a, nomi="Ota")
    bolim = Bolim.objects.create(filial=filial_a, nomi="Bola", ota=ota)
    for _ in range(8):
        av = av_yarat(filial_a, soni="2", turi=tur("T"))
        kartochka(av, oy_qosh(BUGUN, -10), 12)
        qator(av, filial_a, "2", bolim=bolim, mol=ali)
    with django_assert_max_num_queries(8):
        servis.kesim(hisobchi, "bolim", BUGUN)
        servis.mol_qaydnoma(hisobchi)
    with django_assert_max_num_queries(6):
        servis.muddat(hisobchi, BUGUN)


def test_kesim_bolim_yigindisi_1c_summasiga_aniq_teng(client_rol, hisobchi, filial_a):
    av = av_yarat(filial_a, soni="3")  # summa 100, dona 33.333...
    bolimlar = [Bolim.objects.create(filial=filial_a, nomi=n) for n in "XYZ"]
    for b in bolimlar:
        qator(av, filial_a, "1", bolim=b)

    qatorlar = servis.kesim(hisobchi, "bolim", BUGUN)
    assert sorted(q["boshlangich"] for q in qatorlar) == [Decimal("33.33"), Decimal("33.33"), Decimal("33.34")]
    assert servis.jami(qatorlar)["boshlangich"] == Decimal("100.00")
    assert servis.jami(qatorlar)["qoldiq"] == Decimal("100.00")

    client = client_rol("hisobchi")
    javob = client.get(reverse("hisobotlar:kesim"), {"guruh": "bolim"})
    assert javob.context["jami"]["boshlangich"] == Decimal("100.00")
    satrlar = xlsx_qatorlar(client.get(reverse("hisobotlar:kesim"), {"guruh": "bolim", "format": "xlsx"}))
    assert satrlar[-1][0] == "Jami"
    assert Decimal(str(satrlar[-1][2])) == Decimal("100.00")
    assert sum(Decimal(str(q[2])) for q in satrlar[1:-1]) == Decimal("100.00")


def test_kesim_bolim_qoldiq_yigindisi_aniq(hisobchi, filial_a):
    av = av_yarat(filial_a, soni="3")
    kartochka(av, oy_qosh(BUGUN, -7), 24)  # qoldiq 100 * 17/24 = 70.83
    for n in "XYZ":
        qator(av, filial_a, "1", bolim=Bolim.objects.create(filial=filial_a, nomi=n))
    qatorlar = servis.kesim(hisobchi, "bolim", BUGUN)
    assert servis.jami(qatorlar)["qoldiq"] == Decimal("70.83")
    assert servis.jami(qatorlar)["boshlangich"] == Decimal("100.00")


def test_mol_qaydnoma_boshqa_filial_xodimi_404(foydalanuvchi, filial_a, filial_b):
    c = Client()
    c.force_login(foydalanuvchi("hisobchi", [filial_a]))
    begona = xodim(filial=filial_b)
    begona.ism = "Begona Xodim"
    begona.save()
    qator(av_yarat(filial_b), filial_b, "1", mol=begona)
    javob = c.get(reverse("hisobotlar:mol_qaydnoma"), {"xodim": begona.pk})
    assert javob.status_code == 404
    assert "Begona Xodim" not in javob.content.decode()

    oziniki = xodim(filial=filial_a)
    javob = c.get(reverse("hisobotlar:mol_qaydnoma"), {"xodim": oziniki.pk})
    assert javob.status_code == 200 and javob.context["xodim"] == oziniki


def test_mol_qaydnoma_qidiruv_va_mol_tanlash(hisobchi, filial_a):
    """МОЛ qaydnomasi: МОЛ tanlash ro'yxati va matn qidiruvi (МОЛ ismi, AV nomi, inv. raqam); Excel ham filtrlangan."""
    from onec.models import Xodim

    ali = Xodim.objects.create(ref_1c=uuid.uuid4(), ism="Ali Valiyev")
    vali = Xodim.objects.create(ref_1c=uuid.uuid4(), ism="Vali Aliyev")
    stol = av_yarat(filial_a, inv="INV-STOL")
    monitor = av_yarat(filial_a, inv="INV-MON")
    qator(stol, filial_a, "1", mol=ali)
    qator(monitor, filial_a, "1", mol=vali)
    c = Client()
    c.force_login(hisobchi)
    url = reverse("hisobotlar:mol_qaydnoma")

    javob = c.get(url)
    html = javob.content.decode()
    assert 'name="q"' in html and f'<option value="{ali.pk}">Ali Valiyev</option>' in html

    assert [q["av"] for q in c.get(url, {"q": "inv-mon"}).context["qatorlar"]] == [monitor]
    assert [q["av"] for q in c.get(url, {"q": "ali valiyev"}).context["qatorlar"]] == [stol]
    assert [q["av"] for q in c.get(url, {"xodim": vali.pk}).context["qatorlar"]] == [monitor]
    html = c.get(url, {"q": "INV-MON"}).content.decode()
    assert "q=INV-MON&amp;format=xlsx" in html
    kitob = load_workbook(io.BytesIO(c.get(url, {"q": "INV-MON", "format": "xlsx"}).content))
    assert [r[2] for r in kitob.active.iter_rows(min_row=2, values_only=True) if r[2]] == ["INV-MON"]


def test_bolim_qaydnomasi(hisobchi, filial_a, filial_b, foydalanuvchi):
    """Bo'lim qaydnomasi: tanlangan bo'lim va ichki bo'limlari; soni — shu bo'limdagi; МОЛ; qidiruv; Excel; imzo."""
    ali = xodim()
    bino = Bolim.objects.create(filial=filial_a, nomi="Bino")
    xona = Bolim.objects.create(filial=filial_a, nomi="Xona", ota=bino)
    ombor = Bolim.objects.create(filial=filial_a, nomi="Ombor")
    ratsiya = av_yarat(filial_a, inv="INV-R", soni="4")
    qator(ratsiya, filial_a, "1", bolim=xona, mol=ali)
    qator(ratsiya, filial_a, "3", bolim=ombor)
    begona = av_yarat(filial_b, inv="INV-B")
    qator(begona, filial_b, "1")
    c = Client()
    c.force_login(hisobchi)
    url = reverse("hisobotlar:bolim_qaydnoma")
    assert f'href="{url}"' in c.get(reverse("panel:bosh")).content.decode()  # menyuda

    javob = c.get(url, {"filial": filial_a.pk, "bolim": bino.pk})
    [q] = javob.context["qatorlar"]
    assert (q["av"], q["soni"], q["bolim"], q["mol"]) == (ratsiya, Decimal("1"), "Bino / Xona", ali)
    html = javob.content.decode()
    assert "Bo&#x27;lim mas&#x27;uli" in html or "Bo'lim mas'uli" in html
    assert [q["soni"] for q in c.get(url, {"filial": filial_a.pk}).context["qatorlar"]] == [Decimal("1"), Decimal("3")]
    assert c.get(url, {"filial": filial_a.pk, "q": "ombor"}).context["qatorlar"][0]["bolim"] == "Ombor"
    kitob = load_workbook(io.BytesIO(c.get(url, {"filial": filial_a.pk, "bolim": ombor.pk, "format": "xlsx"}).content))
    assert [r[0] for r in kitob.active.iter_rows(min_row=2, values_only=True) if r[0]] == ["Ombor"]
    # boshqa filial bo'limi / ruxsatsiz filial — 404
    cheklangan = Client()
    cheklangan.force_login(foydalanuvchi("hisobchi", [filial_a]))
    assert cheklangan.get(url, {"filial": filial_b.pk}).status_code == 404
    assert c.get(url, {"filial": filial_b.pk, "bolim": ombor.pk}).status_code == 404


def test_qaydnomalarda_xulosa_va_amallar_tepada(hisobchi, filial_a):
    """Uzun ro'yxatda pastga tushmaslik uchun: qatorlar soni, jami va Excel/Chop etish jadval ustida; jadval qotgan."""
    av = av_yarat(filial_a)
    qator(av, filial_a, "1", mol=xodim())
    c = Client()
    c.force_login(hisobchi)
    for url in (reverse("hisobotlar:bolim_qaydnoma") + f"?filial={filial_a.pk}", reverse("hisobotlar:mol_qaydnoma")):
        html = c.get(url).content.decode()
        assert html.index("hisobot-xulosa") < html.index("<table") and "jadval-qotgan" in html
        assert html.index("format=xlsx") < html.index("<table"), url
