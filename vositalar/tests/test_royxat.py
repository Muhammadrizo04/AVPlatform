import io

import pytest
from django.http import QueryDict
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from onec.models import AV
from vositalar.filtrlar import av_filtrla
from vositalar.models import Bolim, Holat, Taqsimot, TaqsimotHarakat
from vositalar.tests.test_bayroqlar import av_yarat, qator, qoldiq

ROYXAT = "vositalar:royxat"


@pytest.fixture
def filial_a(filial_yarat):
    return filial_yarat("A")


@pytest.fixture
def filial_b(filial_yarat):
    return filial_yarat("B")


@pytest.fixture
def client_hisobchi(client_rol):
    return client_rol("hisobchi")


def nomlar(javob):
    return {av.nomi for av in javob.context["sahifa"]}


def xabarlar(javob):
    return " ".join(str(m) for m in javob.context["messages"])


def test_royxat_qidiruv_inv_va_nom(client_hisobchi, filial_a):
    bir = av_yarat(filial_a, inv="ABC-1")
    ikki = av_yarat(filial_a, inv="XYZ-9")
    ikki.nomi = "Noutbuk Lenovo"
    ikki.save()
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"q": "abc"})) == {bir.nomi}
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"q": "lenovo"})) == {"Noutbuk Lenovo"}


@pytest.mark.parametrize("htmx", [False, True])
def test_royxatda_kartochkaga_havola(client_hisobchi, client_rol, filial_a, htmx):
    av = av_yarat(filial_a, inv="ABC-1")
    sarlavha = {"HTTP_HX_REQUEST": "true"} if htmx else {}
    url = reverse("vositalar:kartochka", args=[av.pk])
    for client in (client_hisobchi, client_rol("rahbar")):
        html = client.get(reverse(ROYXAT), **sarlavha).content.decode()  # standart — ixcham jadval
        assert f'<a href="{url}" class="fw-semibold text-decoration-none stretched-link">{av.nomi}</a>' in html
        assert f'<a href="{url}" class="text-body-secondary">ABC-1</a>' in html
        html = client.get(reverse(ROYXAT), {"korinish": "kartochka"}, **sarlavha).content.decode()
        assert f'<a href="{url}" class="fw-semibold text-decoration-none">{av.nomi}</a>' in html
        assert f'<a href="{url}" class="text-body-secondary">Inv. ABC-1</a>' in html


def test_royxat_filtr_bayroq(client_hisobchi, filial_a):
    bosh = av_yarat(filial_a, inv="")
    av_yarat(filial_a, inv="BOR")
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"bayroq": "inv_bosh"})) == {bosh.nomi}


def test_royxat_bayroq_taqsimot_darajasi(client_hisobchi, filial_a):
    av = av_yarat(filial_a)
    qator(av, filial_a)  # bolimsiz
    boshqa = av_yarat(filial_a)
    qator(boshqa, filial_a, bolim=Bolim.objects.create(filial=filial_a, nomi="Ombor"))
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"bayroq": "bolimsiz"})) == {av.nomi}


def test_guruh_filtri(client_hisobchi, filial_a):
    guruh = av_yarat(filial_a, soni="5")
    av_yarat(filial_a, soni="1")
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"guruh": "1"})) == {guruh.nomi}


def test_royxat_filtr_filial_holat_arxiv(client_hisobchi, filial_a, filial_b):
    a = av_yarat(filial_a)
    b = av_yarat(filial_b)
    arxiv = av_yarat(filial_a, soni="0")
    qator(a, filial_a)
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"filial": filial_b.pk})) == {b.nomi}
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"holat": Holat.ISHLATILMOQDA})) == {a.nomi}
    assert nomlar(client_hisobchi.get(reverse(ROYXAT), {"arxiv": "1"})) == {arxiv.nomi}
    assert arxiv.nomi not in nomlar(client_hisobchi.get(reverse(ROYXAT)))


def test_royxat_sahifalash_50(client_hisobchi, filial_a):
    for _ in range(51):
        av_yarat(filial_a)
    javob = client_hisobchi.get(reverse(ROYXAT))
    assert len(javob.context["sahifa"]) == 50
    assert len(client_hisobchi.get(reverse(ROYXAT), {"page": 2}).context["sahifa"]) == 1


def test_royxat_htmx_faqat_jadval(client_hisobchi, filial_a):
    av_yarat(filial_a)
    toliq = client_hisobchi.get(reverse(ROYXAT)).content.decode()
    qism = client_hisobchi.get(reverse(ROYXAT), HTTP_HX_REQUEST="true").content.decode()
    assert "<html" in toliq
    assert "<html" not in qism and "av-jadval" in qism
    assert 'id="av-korsatkichlar" hx-swap-oob="true"' in qism  # plitkalar filtr bilan yangilanadi


def test_royxat_filial_cheklovi(foydalanuvchi, filial_a, filial_b):
    meniki = av_yarat(filial_a)
    boshqa = av_yarat(filial_b)
    arxiv_meniki = av_yarat(filial_a, soni="0")
    arxiv_boshqa = av_yarat(filial_b, soni="0")
    client = Client()
    client.force_login(foydalanuvchi("hisobchi", [filial_a]))
    javob = client.get(reverse(ROYXAT))
    assert nomlar(javob) == {meniki.nomi}
    assert boshqa.nomi not in javob.content.decode()
    # filial parametri bilan ham boshqa filialni ko'rib bo'lmaydi
    assert nomlar(client.get(reverse(ROYXAT), {"filial": filial_b.pk})) == set()
    # arxiv: ruxsat etilgan filial_rekvizit
    arxiv = client.get(reverse(ROYXAT), {"arxiv": "1"})
    assert nomlar(arxiv) == {arxiv_meniki.nomi}
    assert arxiv_boshqa.nomi not in arxiv.content.decode()


def test_av_filtrla_arxiv_cheklanmagan(foydalanuvchi, filial_a, filial_b):
    av_yarat(filial_a, soni="0")
    av_yarat(filial_b, soni="0")
    assert av_filtrla(AV.objects.all(), QueryDict("arxiv=1"), foydalanuvchi("admin")).count() == 2


def test_excel_filial_cheklovi(foydalanuvchi, filial_a, filial_b):
    meniki = av_yarat(filial_a)
    boshqa = av_yarat(filial_b)
    client = Client()
    client.force_login(foydalanuvchi("hisobchi", [filial_a]))
    javob = client.get(reverse("vositalar:royxat_excel"))
    assert javob.status_code == 200
    varaq = load_workbook(io.BytesIO(javob.content)).active
    matn = [str(c) for satr in varaq.iter_rows(values_only=True) for c in satr]
    assert meniki.nomi in matn
    assert boshqa.nomi not in matn


def test_excel_rahbar_va_revizor(client_rol, filial_a):
    av_yarat(filial_a)
    assert client_rol("rahbar").get(reverse("vositalar:royxat_excel")).status_code == 200
    assert client_rol("revizor").get(reverse("vositalar:royxat_excel")).status_code == 403


@pytest.fixture
def bolim(filial_a):
    return Bolim.objects.create(filial=filial_a, nomi="Ombor")


def test_ommaviy_kochirish_faqat_1_donali(client_hisobchi, filial_a, bolim):
    yakka = av_yarat(filial_a, soni="1")
    qator(yakka, filial_a)
    guruh = av_yarat(filial_a, soni="3")
    qator(guruh, filial_a, soni="3")
    javob = client_hisobchi.post(
        reverse("vositalar:ommaviy"),
        {"av": [yakka.pk, guruh.pk], "amal": "kochirish", "bolim": bolim.pk},
        follow=True,
    )
    assert Taqsimot.objects.get(av=yakka).bolim == bolim
    assert Taqsimot.objects.get(av=guruh).bolim is None
    assert TaqsimotHarakat.objects.filter(av=guruh).count() == 0
    assert guruh.nomi in xabarlar(javob)


def test_ommaviy_holat(client_hisobchi, filial_a):
    av = av_yarat(filial_a)
    qator(av, filial_a)
    client_hisobchi.post(reverse("vositalar:ommaviy"), {"av": [av.pk], "amal": "holat", "holat": Holat.NOSOZ})
    assert Taqsimot.objects.get(av=av).holat == Holat.NOSOZ


def test_ommaviy_amal_xatosi_500_emas(client_hisobchi, filial_a):
    av = av_yarat(filial_a)
    qator(av, filial_a)  # holat allaqachon ishlatilmoqda -> AmalXato
    javob = client_hisobchi.post(
        reverse("vositalar:ommaviy"), {"av": [av.pk], "amal": "holat", "holat": Holat.ISHLATILMOQDA}, follow=True
    )
    assert javob.status_code == 200
    assert av.nomi in xabarlar(javob)


def test_ommaviy_boshqa_filial_av_tegilmaydi(foydalanuvchi, filial_a, filial_b):
    boshqa = av_yarat(filial_b)
    qator(boshqa, filial_b)
    bolim_a = Bolim.objects.create(filial=filial_a, nomi="A ombor")
    client = Client()
    client.force_login(foydalanuvchi("hisobchi", [filial_a]))
    javob = client.post(
        reverse("vositalar:ommaviy"), {"av": [boshqa.pk], "amal": "kochirish", "bolim": bolim_a.pk}, follow=True
    )
    assert Taqsimot.objects.get(av=boshqa).bolim is None
    # darvoza AV'ni umuman ko'rmaydi: u qayta ishlanmaydi (AmalXato xabarida ham tilga olinmaydi)
    assert boshqa.nomi not in xabarlar(javob)


def test_royxat_ruxsatsiz_filial_parametri_bosh(foydalanuvchi, filial_a, filial_b):
    ikkalasida = av_yarat(filial_a)
    qoldiq(ikkalasida, filial_b, "1")
    client = Client()
    client.force_login(foydalanuvchi("hisobchi", [filial_a]))
    assert nomlar(client.get(reverse(ROYXAT), {"filial": filial_a.pk})) == {ikkalasida.nomi}
    assert nomlar(client.get(reverse(ROYXAT), {"filial": filial_b.pk})) == set()


def test_htmx_qismida_excel_havolasi_joriy_filtr_bilan(client_hisobchi, filial_a):
    av_yarat(filial_a, inv="ABC")
    qism = client_hisobchi.get(reverse(ROYXAT), {"q": "abc"}, HTTP_HX_REQUEST="true").content.decode()
    assert f"{reverse('vositalar:royxat_excel')}?q=abc" in qism


def test_royxat_mol_va_bayroq_boshqaruvlari(client_hisobchi, filial_a):
    from vositalar.tests.test_bayroqlar import xodim

    x = xodim(filial=filial_a)
    matn = client_hisobchi.get(reverse(ROYXAT)).content.decode()
    assert 'name="mol_ism"' in matn and f'<option value="{x.ism}">' in matn  # МОЛ ismi bo'yicha qidiruv (datalist)
    assert 'name="bayroq"' in matn and 'value="inv_bosh"' in matn


def test_rahbar_ommaviy_amal_403(client_rol, filial_a, bolim):
    av = av_yarat(filial_a)
    qator(av, filial_a)
    javob = client_rol("rahbar").post(
        reverse("vositalar:ommaviy"), {"av": [av.pk], "amal": "kochirish", "bolim": bolim.pk}
    )
    assert javob.status_code == 403
    assert Taqsimot.objects.get(av=av).bolim is None


def test_ommaviy_faqat_post(client_hisobchi):
    assert client_hisobchi.get(reverse("vositalar:ommaviy")).status_code == 405


def test_royxat_n_plus_1_yoq(client_hisobchi, filial_a, django_assert_max_num_queries):
    for _ in range(3):
        av_yarat(filial_a)
    # Qator joylashuvi (bitta so'rov) va ko'rsatkichlar qatori (3 ta bayroq soni) — AV soniga bog'liq emas.
    with django_assert_max_num_queries(36):
        client_hisobchi.get(reverse(ROYXAT))
    for _ in range(20):
        av_yarat(filial_a)
    with django_assert_max_num_queries(36):
        client_hisobchi.get(reverse(ROYXAT))


def test_korsatkichlar_filtrga_mos(client_hisobchi, filial_a, filial_b):
    qator(av_yarat(filial_a, inv="A1"), filial_a)  # МОЛ'siz va bo'limsiz
    qator(av_yarat(filial_b, inv="B1"), filial_b)
    qator(av_yarat(filial_b, inv="B2"), filial_b)

    hammasi = client_hisobchi.get(reverse(ROYXAT)).context["korsatkichlar"]
    assert (hammasi["molsiz"], hammasi["bolimsiz"]) == (3, 3)
    javob = client_hisobchi.get(reverse(ROYXAT), {"filial": filial_a.pk, "tartib": "-soni"})
    assert (javob.context["korsatkichlar"]["molsiz"], javob.context["korsatkichlar"]["bolimsiz"]) == (1, 1)
    html = javob.content.decode()
    # plitka joriy filtrni saqlab bayroq qo'shadi; tozalash tugmasi faol filtrlar soni bilan
    assert f"?filial={filial_a.pk}&amp;tartib=-soni&amp;bayroq=molsiz" in html
    assert "Filtrlarni tozalash" in html and 'badge text-bg-secondary ms-1">1<' in html
    assert "Filtrlarni tozalash" not in client_hisobchi.get(reverse(ROYXAT)).content.decode()


def test_amallar_paneli_va_hammasini_tanlash(client_hisobchi, client_rol, filial_a):
    av_yarat(filial_a)
    html = client_hisobchi.get(reverse(ROYXAT)).content.decode()
    assert '<div class="amal-panel" id="amal-panel" hidden>' in html and 'type="checkbox" data-hammasi' in html
    # rahbar — taqsimot va yorliq huquqi yo'q: tanlov va panel chiqmaydi
    html = client_rol("rahbar").get(reverse(ROYXAT)).content.decode()
    assert '<div class="amal-panel"' not in html and 'type="checkbox" data-hammasi' not in html
